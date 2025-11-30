#!/usr/bin/env python3
"""
Test script to process the example_2.docx file and check chapter parsing
"""

import sys
import os
from pathlib import Path

# Add the backend directory to the path
sys.path.append(os.path.dirname(__file__))

def test_with_example_docx():
    """Test chapter parsing with the actual example_2.docx file"""
    try:
        from docx import Document
        import io
        
        # Load the example document
        example_path = Path(__file__).parent / "templates" / "example_2.docx"
        
        if not example_path.exists():
            print(f"❌ Example file not found: {example_path}")
            return False
        
        print(f"📄 Loading document: {example_path}")
        
        # Read the .docx file
        doc = Document(example_path)
        
        # Extract text from all paragraphs
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text.strip())
        
        document_content = '\n\n'.join(paragraphs)
        
        print(f"📊 Document Stats:")
        print(f"   - Total paragraphs: {len(paragraphs)}")
        print(f"   - Total characters: {len(document_content)}")
        print(f"   - Total words: {len(document_content.split())}")
        
        # Show first few paragraphs to understand structure
        print(f"\n📖 First 10 paragraphs:")
        for i, para in enumerate(paragraphs[:10]):
            print(f"   {i+1}: {para[:100]}{'...' if len(para) > 100 else ''}")
        
        # Now test our chapter parsing
        chapters = parse_docx_into_chapters(document_content)
        
        print(f"\n🔍 CHAPTER PARSING RESULTS:")
        print(f"   Found {len(chapters)} chapters")
        print("="*80)
        
        for i, chapter in enumerate(chapters):
            print(f"\n📖 Chapter {chapter['chapter_number']}: {chapter['title']}")
            print(f"   Content length: {len(chapter['content'])} characters")
            print(f"   Word count: {len(chapter['content'].split())} words")
            print(f"   Image count: {chapter['image_count']}")
            
            # Show content preview (first 300 chars)
            content_preview = chapter['content'][:300].replace('\n', ' ')
            print(f"   Preview: {content_preview}{'...' if len(chapter['content']) > 300 else ''}")
            
            # Show image placeholders if any
            if chapter['image_count'] > 0:
                import re
                image_pattern = r'\[IMAGE:\d+_\d+[^\]]*\]'
                images = re.findall(image_pattern, chapter['content'])
                print(f"   Images found:")
                for img in images:
                    print(f"     - {img}")
        
        return True
        
    except ImportError:
        print("❌ python-docx not installed. Install with: pip install python-docx")
        return False
    except Exception as e:
        print(f"❌ Error processing document: {e}")
        return False

def parse_docx_into_chapters(content: str) -> list:
    """
    Parse document content into chapters using regex pattern 'CHAPTER *'
    (Copy of the function from app.py for testing)
    """
    import re
    
    print(f"\n🔍 Starting chapter parsing. Content length: {len(content)} characters")
    
    # Enhanced chapter detection patterns
    chapter_patterns = [
        r'(?i)CHAPTER\s+(\d+)[:\-\.\s]*([^\n]*)',  # CHAPTER 1: Title or CHAPTER 1 Title
        r'(?i)CHAPTER\s*[:\-\.]?\s*(\d+)[:\-\.\s]*([^\n]*)',  # CHAPTER: 1 Title
        r'(?i)(\d+)\.\s*CHAPTER[:\-\.\s]*([^\n]*)',  # 1. CHAPTER: Title
        r'(?i)CHAPTER\s*([IVXLCDM]+)[:\-\.\s]*([^\n]*)',  # CHAPTER I: Title (Roman numerals)
    ]
    
    # Find all chapter matches with their positions
    chapter_matches = []
    
    for pattern_idx, pattern in enumerate(chapter_patterns):
        matches = re.finditer(pattern, content, re.MULTILINE)
        for match in matches:
            # Extract chapter number and title
            groups = match.groups()
            if len(groups) >= 2:
                chapter_num_str = groups[0].strip()
                chapter_title = groups[1].strip()
                
                # Convert Roman numerals to integers if needed
                if re.match(r'^[IVXLCDM]+$', chapter_num_str.upper()):
                    chapter_num = roman_to_int(chapter_num_str.upper())
                else:
                    try:
                        chapter_num = int(chapter_num_str)
                    except ValueError:
                        continue
                
                chapter_matches.append({
                    'start_pos': match.start(),
                    'end_pos': match.end(),
                    'chapter_number': chapter_num,
                    'title': chapter_title if chapter_title else f"Chapter {chapter_num}",
                    'match_text': match.group(0),
                    'pattern_used': pattern_idx + 1
                })
    
    # If no specific chapter patterns found, try a simpler approach
    if not chapter_matches:
        print("⚠️  No specific chapter patterns found, trying simpler approach")
        simple_pattern = r'(?i)(CHAPTER[^\n]*)'
        matches = re.finditer(simple_pattern, content, re.MULTILINE)
        for i, match in enumerate(matches):
            chapter_matches.append({
                'start_pos': match.start(),
                'end_pos': match.end(),
                'chapter_number': i + 1,
                'title': match.group(1).strip(),
                'match_text': match.group(0),
                'pattern_used': 'simple'
            })
    
    # Sort chapters by their position in the document
    chapter_matches.sort(key=lambda x: x['start_pos'])
    
    print(f"✅ Found {len(chapter_matches)} chapter markers:")
    for match in chapter_matches:
        print(f"   - Chapter {match['chapter_number']}: '{match['title']}' (pattern {match['pattern_used']}) at pos {match['start_pos']}")
    
    if not chapter_matches:
        # No chapters found, treat entire document as single chapter
        print("⚠️  No chapters detected, treating as single document")
        return [{
            'chapter_number': 1,
            'title': 'Document Content',
            'content': process_chapter_content(content, 1),
            'image_count': count_images_in_content(content)
        }]
    
    parsed_chapters = []
    
    for i, chapter_match in enumerate(chapter_matches):
        # Determine chapter content boundaries
        content_start = chapter_match['end_pos']
        
        # Content ends at the start of next chapter or end of document
        if i + 1 < len(chapter_matches):
            content_end = chapter_matches[i + 1]['start_pos']
        else:
            content_end = len(content)
        
        # Extract chapter content
        chapter_content = content[content_start:content_end].strip()
        
        # Process the content (handle images, clean formatting)
        processed_content = process_chapter_content(chapter_content, chapter_match['chapter_number'])
        
        # Count images in this chapter
        image_count = count_images_in_content(processed_content)
        
        parsed_chapters.append({
            'chapter_number': chapter_match['chapter_number'],
            'title': clean_chapter_title(chapter_match['title']),
            'content': processed_content,
            'image_count': image_count
        })
        
        print(f"✅ Parsed Chapter {chapter_match['chapter_number']}: {chapter_match['title'][:50]}... ({len(processed_content)} chars, {image_count} images)")
    
    # Sort chapters by chapter number to ensure proper order
    parsed_chapters.sort(key=lambda x: x['chapter_number'])
    
    print(f"🎉 Successfully parsed {len(parsed_chapters)} chapters in correct order")
    return parsed_chapters

def roman_to_int(roman: str) -> int:
    """Convert Roman numeral to integer"""
    roman_numerals = {'I': 1, 'V': 5, 'X': 10, 'L': 50, 'C': 100, 'D': 500, 'M': 1000}
    result = 0
    prev_value = 0
    
    for char in reversed(roman):
        value = roman_numerals.get(char, 0)
        if value < prev_value:
            result -= value
        else:
            result += value
        prev_value = value
    
    return result

def clean_chapter_title(title: str) -> str:
    """Clean and format chapter title"""
    if not title:
        return "Untitled Chapter"
    
    # Remove common prefixes and clean up
    title = re.sub(r'^[:\-\.\s]+', '', title).strip()
    title = re.sub(r'[:\-\.\s]+$', '', title).strip()
    
    # Remove chapter number if it appears at the start
    title = re.sub(r'^\d+[\.\:\-\s]*', '', title).strip()
    
    if not title:
        return "Untitled Chapter"
    
    return title

def process_chapter_content(content: str, chapter_num: int) -> str:
    """Process chapter content to handle images and formatting"""
    if not content:
        return ""
    
    import re
    
    # Replace image references with standardized placeholders
    image_counter = 1
    
    # Enhanced image detection patterns
    image_patterns = [
        r'(?i)\[image[:\s]*([^\]]*)\]',  # [image: description]
        r'(?i)\[fig[ure]*[:\s]*([^\]]*)\]',  # [figure: description]  
        r'(?i)\[insert\s+image[:\s]*([^\]]*)\]',  # [insert image: description]
        r'(?i)<image[^>]*>([^<]*)</image>',  # <image>description</image>
        r'(?i)\{image[:\s]*([^\}]*)\}',  # {image: description}
        r'(?i)image\s*\d*[:\s]*([^\n]*)',  # image: description
        r'(?i)figure\s*\d*[:\s]*([^\n]*)',  # figure: description
    ]
    
    for pattern in image_patterns:
        def replace_image(match):
            nonlocal image_counter
            caption = match.group(1).strip() if match.group(1) else f"Image {image_counter}"
            # Clean up caption
            caption = re.sub(r'^[:\-\.\s]+', '', caption).strip()
            if not caption:
                caption = f"Chapter {chapter_num} Image {image_counter}"
            
            placeholder = f'[IMAGE:{chapter_num}_{image_counter} caption="{caption}"]'
            image_counter += 1
            return placeholder
        
        content = re.sub(pattern, replace_image, content)
    
    # Clean up extra whitespace and formatting
    content = re.sub(r'\n\s*\n\s*\n', '\n\n', content)  # Multiple newlines to double
    content = re.sub(r'[ \t]+', ' ', content)  # Multiple spaces to single
    content = content.strip()
    
    return content

def count_images_in_content(content: str) -> int:
    """Count image placeholders in content"""
    import re
    pattern = r'\[IMAGE:\d+_\d+[^\]]*\]'
    matches = re.findall(pattern, content)
    return len(matches)

if __name__ == "__main__":
    print("🚀 Testing Chapter Parsing with example_2.docx")
    print("=" * 80)
    
    # Add missing import
    import re
    
    success = test_with_example_docx()
    
    if success:
        print("\n🎉 Test completed successfully!")
    else:
        print("\n❌ Test failed. Check the errors above.")